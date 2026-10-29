"""Local end-to-end test of import_questions.py on a throw-away PostgreSQL (pip package pgserver).

Flow: stub + reconstructed schema -> approve some rows in a COPY of the Excel -> generate SQL -> run it twice ->
check upsert, topics, get_daily_set, answer checking, RLS -> negative cases.
Never touches Supabase or out/kimya_dim_yoxlama.xlsx.

Usage: python test_import_local.py
"""
import re
import shutil
import sys
import tempfile
import uuid
from pathlib import Path

import pgserver
import psycopg
from openpyxl import load_workbook

import import_questions as imp

HERE = Path(__file__).parent
DB_DIR = HERE.parent.parent / "db" / "test"
SOURCE = HERE / "out" / "kimya_dim_yoxlama.xlsx"

failures = []


def check(name, condition, detail=""):
    print(("PASS " if condition else "FAIL ") + name + (f"  [{detail}]" if detail and not condition else ""))
    if not condition:
        failures.append(name)


def make_db(srv, name, schema_patch=None):
    admin = psycopg.connect(srv.get_uri(), autocommit=True)
    admin.execute(f'drop database if exists "{name}"')
    admin.execute(f'create database "{name}"')
    admin.close()
    conn = psycopg.connect(srv.get_uri(database=name), autocommit=True)
    conn.execute((DB_DIR / "supabase_stub.sql").read_text(encoding="utf-8"))
    schema = (DB_DIR / "schema_reconstructed.sql").read_text(encoding="utf-8")
    conn.execute(schema_patch(schema) if schema_patch else schema)
    conn.execute(
        "insert into subjects (slug, name, sort_order, is_active) values "
        "('riyaziyyat','Riyaziyyat',1,false), ('azerbaycan-dili','Azərbaycan dili',2,false), "
        "('ingilis-dili','İngilis dili',3,false), ('kimya','Kimya',1,true)"
    )
    return conn


def approve(path, rows):
    """rows: {kod: (çətinlik, aktiv)} written into a copy of the Excel."""
    wb = load_workbook(path)
    ws = wb["suallar"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    for r in range(2, ws.max_row + 1):
        kod = ws.cell(r, col["kod"]).value
        if kod in rows:
            diff, active = rows[kod]
            ws.cell(r, col["çətinlik"]).value = diff
            ws.cell(r, col["aktiv"]).value = active
    wb.save(path)


def reset_decisions(path):
    """The real Excel already holds approvals; tests must start from a clean slate."""
    wb = load_workbook(path)
    ws = wb["suallar"]
    header = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, header.index("aktiv") + 1).value = "xeyr"
        ws.cell(r, header.index("çətinlik") + 1).value = None
    wb.save(path)


def pick_codes():
    ws = load_workbook(SOURCE)["suallar"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    clean = [r["kod"] for r in rows if not r["yoxlanmalı"] and r["tip"] == "qapali"][:8]
    numeric = next(r["kod"] for r in rows if r["tip"] == "aciq_eded")
    image = next(r["kod"] for r in rows if r["tip"] == "qapali" and r["şəkil"] and r["kod"] not in clean)
    written = next(r["kod"] for r in rows if r["tip"] == "yazili")
    no_diff = next(r["kod"] for r in rows if r["tip"] == "qapali" and r["kod"] not in clean + [image])
    return clean, numeric, image, written, no_diff


def main():
    clean, numeric, image, written, no_diff = pick_codes()
    approved = clean + [numeric, image]
    work = Path(tempfile.mkdtemp())
    excel = work / "test.xlsx"
    shutil.copy(SOURCE, excel)
    reset_decisions(excel)
    approve(excel, {c: (3, "bəli") for c in approved} | {written: (2, "bəli"), no_diff: ("", "bəli")})

    # ---- generate SQL
    sql_path = work / "import.sql"
    code = imp.main(["--file", str(excel), "--out", str(sql_path)])
    check("import script succeeds", code == 0 and sql_path.exists())
    sql = sql_path.read_text(encoding="utf-8")
    check("only approved qapali/aciq_eded rows with çətinlik are in the SQL",
          all(f"'{c}'" in sql for c in approved) and f"'{written}'" not in sql.split("-- 1)")[1].split("-- 4)")[0]
          and f"'{no_diff}'" not in sql.split("-- 3)")[1].split("-- 4)")[0])
    check("no secrets or service keys in the SQL", not re.search(r"service_role|eyJ|password", sql, re.I))

    srv = pgserver.get_server(Path(tempfile.mkdtemp()) / "pg")
    try:
        migration_003 = (HERE.parent.parent / "db" / "migrations" / "003_solution_image.sql").read_text(encoding="utf-8")
        db0 = make_db(srv, "kimya_no003")
        try:
            db0.execute(sql)
            check("import without migration 003 stops at the pre-flight", False)
        except psycopg.errors.RaiseException as e:
            check("import without migration 003 stops at the pre-flight", "003" in str(e))
        db0.execute("rollback")
        check("nothing written without migration 003", db0.execute("select count(*) from questions").fetchone()[0] == 0)

        db = make_db(srv, "kimya_test")
        before = db.execute("select count(*) from get_daily_set(1::smallint, 5)").fetchone()[0]
        db.execute(migration_003)
        db.execute(migration_003)
        check("migration 003 is idempotent and get_daily_set still works after it",
              db.execute("select count(*) from get_daily_set(1::smallint, 5)").fetchone()[0] == before == 0)
        check("migration 003 adds solution_image_path",
              db.execute("select count(*) from information_schema.columns where table_name = 'questions' "
                         "and column_name = 'solution_image_path'").fetchone()[0] == 1)
        topics_before = db.execute("select count(*) from topics").fetchone()[0]
        db.execute(sql)
        n_q = db.execute("select count(*) from questions").fetchone()[0]
        check("questions loaded", n_q == len(approved), f"{n_q} vs {len(approved)}")
        check("sections created (parent_id null) and topics under them",
              db.execute("select count(*) from topics where parent_id is null and name like '%kimya'").fetchone()[0] >= 1
              and db.execute("select count(*) from topics t join topics p on p.id = t.parent_id").fetchone()[0] >= 1)
        check("source/source_ref/exam level/groups stored",
              db.execute("select count(*) from questions where source = 'DİM' and source_ref like '%variantı%' "
                         "and exam_levels = '{qebul}' and groups <> '{}' and is_active").fetchone()[0] == n_q)
        check("image_path stored for the image question",
              db.execute("select image_path from questions where code = %s", (image,)).fetchone()[0] == f"{image}.png")
        check("every question has its explanation image path stored",
              db.execute("select count(*) from questions where solution_image_path like 'KIM-%_izah.png'").fetchone()[0] == n_q)
        check("aciq_eded has no options, qapali has A-E",
              db.execute("select options is null from questions where code = %s", (numeric,)).fetchone()[0]
              and db.execute("select count(*) from jsonb_object_keys((select options from questions where code = %s))",
                             (clean[0],)).fetchone()[0] == 5)

        # ---- re-run: no duplicates
        db.execute(sql)
        check("re-running creates no duplicate questions",
              db.execute("select count(*) from questions").fetchone()[0] == n_q)
        topics_after = db.execute("select count(*) from topics").fetchone()[0]
        db.execute(sql)
        check("re-running creates no duplicate topics",
              db.execute("select count(*) from topics").fetchone()[0] == topics_after and topics_after > topics_before)

        # ---- change a value in Excel -> upsert updates it
        approve(excel, {clean[0]: (5, "bəli")})
        imp.main(["--file", str(excel), "--out", str(sql_path)])
        db.execute(sql_path.read_text(encoding="utf-8"))
        check("upsert updates changed fields",
              db.execute("select difficulty from questions where code = %s", (clean[0],)).fetchone()[0] == 5
              and db.execute("select count(*) from questions").fetchone()[0] == n_q)

        # ---- the app's view: authenticated user
        uid = str(uuid.uuid4())
        db.execute("insert into auth.users (id, email) values (%s, 'a@b.az')", (uid,))
        db.execute("insert into profiles (id) values (%s)", (uid,))
        db.execute("select set_config('request.jwt.claim.sub', %s, false)", (uid,))
        db.execute("set role authenticated")
        subjects = [r[0] for r in db.execute("select name from subjects").fetchall()]
        check("only Kimya is visible", subjects == ["Kimya"], str(subjects))
        db.execute("reset role")
        kimya_id = db.execute("select id from subjects where slug = 'kimya'").fetchone()[0]
        db.execute("set role authenticated")
        daily = db.execute("select code, type, correct_answer, image_path, solution_image_path "
                           "from get_daily_set(%s::smallint, 50)", (kimya_id,)).fetchall()
        check("get_daily_set returns the explanation image path", all(r[4] for r in daily))
        check("get_daily_set returns all loaded chemistry questions", len(daily) == n_q, f"{len(daily)} vs {n_q}")

        by_code = {r[0]: r for r in daily}
        right = by_code[clean[0]][2]
        wrong = next(x for x in "ABCDE" if x != right)
        ids = {c: db.execute("select id from questions where code = %s", (c,)).fetchone()[0] for c in (clean[0], clean[1], numeric)}
        db.execute("insert into answers (user_id, question_id, given_answer) values (%s, %s, %s)", (uid, ids[clean[0]], right.lower()))
        db.execute("insert into answers (user_id, question_id, given_answer) values (%s, %s, %s)", (uid, ids[clean[1]], "Z"))
        num_right = by_code[numeric][2]
        db.execute("insert into answers (user_id, question_id, given_answer) values (%s, %s, %s)",
                   (uid, ids[numeric], num_right.replace(",", ".")))
        res = dict(db.execute("select q.code, a.is_correct from answers a join questions q on q.id = a.question_id").fetchall())
        check("server marks correct qapali answer (case-insensitive)", res[clean[0]] is True)
        check("server marks wrong qapali answer", res[clean[1]] is False)
        check("server accepts numeric answer with dot or comma", res[numeric] is True, f"given {num_right!r}")
        check("wrong answer goes to review_queue",
              db.execute("select count(*) from review_queue").fetchone()[0] == 1)
        check("topic_mastery shows the unified topic name",
              db.execute("select count(*) from topic_mastery where topic_name = "
                         "(select t.name from questions q join topics t on t.id = q.topic_id where q.code = %s)",
                         (clean[0],)).fetchone()[0] == 1)
        db.execute("reset role")

        # ---- migration 002: private bucket, read-only for signed-in users, idempotent
        bucket_sql = (HERE.parent.parent / "db" / "migrations" / "002_question_images_bucket.sql").read_text(encoding="utf-8")
        db.execute(bucket_sql)
        db.execute(bucket_sql)
        check("bucket created once, private, size/mime limited",
              db.execute("select count(*) from storage.buckets where id = 'question-images' and not public "
                         "and file_size_limit = 2097152").fetchone()[0] == 1)
        check("exactly one read policy on storage.objects",
              db.execute("select count(*) from pg_policies where schemaname = 'storage' "
                         "and policyname = 'question images read'").fetchone()[0] == 1)
        db.execute("insert into storage.objects (bucket_id, name) values ('question-images', 'KIM-0001.png')")
        db.execute("set role authenticated")
        check("signed-in user can read images",
              db.execute("select count(*) from storage.objects where bucket_id = 'question-images'").fetchone()[0] == 1)
        try:
            db.execute("insert into storage.objects (bucket_id, name) values ('question-images', 'hack.png')")
            check("signed-in user cannot upload", False)
        except psycopg.errors.InsufficientPrivilege:
            check("signed-in user cannot upload", True)
        db.execute("reset role")
        check("migration 002 leaves subjects.is_active untouched",
              db.execute("select count(*) from subjects where is_active").fetchone()[0] == 1)

        # ---- negative: pre-flight stops when the enum label is missing (nothing is written)
        db2 = make_db(srv, "kimya_noenum", lambda s: s.replace("('buraxilis', 'qebul')", "('buraxilis')"))
        try:
            db2.execute(sql)
            check("pre-flight stops when exam_level has no 'qebul'", False)
        except psycopg.errors.RaiseException as e:
            check("pre-flight stops when exam_level has no 'qebul'", "qebul" in str(e))
        db2.execute("rollback")  # the script's own BEGIN was left open by the exception
        check("nothing was written after the failed pre-flight",
              db2.execute("select count(*) from questions").fetchone()[0] == 0)
    finally:
        srv.cleanup()

    # ---- negative: invalid approved rows -> errors with row numbers, no SQL file
    bad_excel, bad_sql = work / "bad.xlsx", work / "bad.sql"
    shutil.copy(SOURCE, bad_excel)
    reset_decisions(bad_excel)
    approve(bad_excel, {clean[0]: (9, "bəli"), clean[1]: (3, "bəli")})
    wb = load_workbook(bad_excel)
    ws = wb["suallar"]
    header = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == clean[1]:
            ws.cell(r, header.index("düzgün cavab") + 1).value = "Z"
    wb.save(bad_excel)
    import io
    import contextlib
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = imp.main(["--file", str(bad_excel), "--out", str(bad_sql)])
    out = buf.getvalue()
    check("invalid rows: exit code 1, no SQL written, errors list row numbers",
          code == 1 and not bad_sql.exists() and re.search(r"sətir \d+ \(" + clean[0], out) is not None
          and re.search(r"sətir \d+ \(" + clean[1], out) is not None, out[-300:])

    shutil.rmtree(work, ignore_errors=True)
    print(f"\n{'ALL PASSED' if not failures else 'FAILED: ' + ', '.join(failures)}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
