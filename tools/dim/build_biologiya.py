"""Builds out/biologiya_dim_yoxlama.xlsx from out/parsed_biologiya.json (run extract_biologiya.py first).

Same columns and review flow as chemistry. Identical questions found in several exams are merged into one row
(their groups, years and places are combined); questions that are only similar get a "oxşar" note for the teacher.

Usage: python build_biologiya.py [--force]
"""
import collections
import difflib
import json
import re
import sys

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from build_xlsx import COLUMNS, YELLOW, build_row
from chem_text import MARKS, SUB_C, SUB_O, SUP_C, SUP_O
from extract_biologiya import BIO_FILES
from extract_kimya import IMAGES, RAW, collapse, crop_figure
from subjects import SUBJECTS
from topic_map_biologiya import map_topic

CFG = SUBJECTS["biologiya"]
SIMILAR = 0.9
_SUB = str.maketrans("0123456789+-", "₀₁₂₃₄₅₆₇₈₉₊₋")
_SUP = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")


def plain(raw, _flags):
    """PDF text -> plain text; sub/superscripts become unicode (CO₂, Ca²⁺)."""
    text = collapse(raw)
    text = re.sub(f"{SUB_O}(.*?){SUB_C}", lambda m: m.group(1).translate(_SUB), text)
    text = re.sub(f"{SUP_O}(.*?){SUP_C}", lambda m: m.group(1).translate(_SUP), text)
    return re.sub(f"[{MARKS}]", "", text)


def merge(rows):
    """Exact duplicates are merged; returns (unique rows, number of merged duplicates)."""
    merged, seen, dropped = [], {}, 0
    for r in rows:
        b = build_row(r, plain)
        first = seen.get(b["dedup"])
        if first is None:
            seen[b["dedup"]] = b
            merged.append(b)
            continue
        dropped += 1
        for key in ("groups", "years", "places"):
            first[key] += b[key]
    return merged, dropped


def mark_similar(rows):
    """Adds an 'oxşar' reason to rows whose text is very close but not identical (e.g. a changed number)."""
    pairs = 0
    for i, a in enumerate(rows):
        for j in range(i + 1, len(rows)):
            b = rows[j]
            sm = difflib.SequenceMatcher(None, a["dedup"], b["dedup"])
            if sm.real_quick_ratio() < SIMILAR or sm.quick_ratio() < SIMILAR or sm.ratio() < SIMILAR:
                continue
            a["similar"].append(j)
            b["similar"].append(i)
            pairs += 1
    return pairs


def main():
    target = CFG["excel"]
    if target.exists() and "--force" not in sys.argv:
        sys.exit(
            f"{target.name} already exists. Codes (BIO-xxxx) are keys in the database and the teacher edits this file: "
            "rebuilding would renumber them. Pass --force only before any review or upload."
        )
    rows = json.loads((target.parent / "parsed_biologiya.json").read_text(encoding="utf-8"))
    per_file = collections.Counter(r["file"] for r in rows)
    merged, dropped = merge(rows)
    for b in merged:
        b["similar"] = []
    pairs = mark_similar(merged)
    codes = [f"{CFG['prefix']}-{n:04d}" for n in range(1, len(merged) + 1)]

    for old in IMAGES.glob(f"{CFG['prefix']}-*.png"):
        old.unlink()
    docs = {stem: pymupdf.open(RAW / f"{stem}.pdf") for stem, *_ in BIO_FILES}
    wb = Workbook()
    ws = wb.active
    ws.title = "suallar"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    flagged_by_file, topic_counts = collections.Counter(), collections.Counter()
    for kod, b in zip(codes, merged):
        image = expl_image = ""
        if b["crop"]:
            crop_figure(docs[b["crop"][0]], b["crop"][1], kod)
            image = f"{kod}.png"
        if b["expl_crop"]:
            expl_image = f"{kod}_izah.png"
            crop_figure(docs[b["file"]], b["expl_crop"], f"{kod}_izah", pad=0)
        section, topic, note = map_topic(b["topic"])
        reasons = list(b["reasons"])
        if note:
            reasons.append(note)
        if b["similar"]:
            reasons.append("oxşar sual: " + ", ".join(codes[j] for j in b["similar"]))
        flagged = bool(reasons)
        if flagged:
            flagged_by_file[b["file"]] += 1
        topic_counts[(section, topic)] += 1
        opts = b["options"]
        ws.append([
            kod, CFG["name"], section, topic, b["topic"], b["type"], b["stem"],
            opts.get("A", ""), opts.get("B", ""), opts.get("C", ""), opts.get("D", ""), opts.get("E", ""),
            b["correct"], "", b["solution"], image, "qebul",
            ",".join(dict.fromkeys(b["groups"])), ",".join(str(y) for y in dict.fromkeys(b["years"])),
            "DİM", "; ".join(b["places"]), "xeyr", "bəli" if flagged else "", "; ".join(reasons), b["sinif"],
            expl_image,
        ])
        if flagged:
            for c in ws[ws.max_row]:
                c.fill = YELLOW
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVWXYZ", (10, 10, 22, 34, 34, 10, 60, 24, 24, 24, 24, 24, 12, 9, 50, 14, 9, 8, 8, 6, 34, 7, 11, 40, 7, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    s = wb.create_sheet("xülasə")
    s.append(["Fayl", "İmtahan", "Çıxarılan sual", "Bu faylda işarələnmiş (dublikatsız)"])
    for stem, exam, _g, _y in BIO_FILES:
        s.append([stem, exam, per_file[stem], flagged_by_file[stem]])
    s.append([])
    s.append(["Cəmi çıxarılan", "", sum(per_file.values()), ""])
    s.append(["Dublikatsız sual", "", len(merged), sum(flagged_by_file.values())])
    s.append([])
    s.append(["Bölmə", "Mövzu", "Sual sayı"])
    order = {name: n for n, name in enumerate(CFG["sections"])}
    for (sec, topic), c in sorted(topic_counts.items(), key=lambda kv: (order.get(kv[0][0], 9), kv[0][1])):
        s.append([sec, topic, c])
    s.column_dimensions["A"].width = 40
    s.column_dimensions["B"].width = 50
    s.column_dimensions["C"].width = 16
    s.column_dimensions["D"].width = 34
    wb.save(target)
    print(f"questions {len(merged)} (from {sum(per_file.values())}, {dropped} exact duplicates merged, {pairs} similar pairs)")
    print(f"flagged {sum(flagged_by_file.values())}; unmapped topics: {topic_counts[('', '')]}")


if __name__ == "__main__":
    main()
