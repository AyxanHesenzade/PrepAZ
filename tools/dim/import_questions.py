"""Generates out/import_kimya.sql from the reviewed Excel. The owner runs that SQL in the Supabase SQL Editor.

Only rows with aktiv = bəli, tip qapali/aciq_eded and a filled çətinlik are loaded; the rest are skipped
and counted. Everything is validated first: one error means no SQL file at all.
The SQL is one transaction and safe to run again: topics are created only if missing, questions are upserted by code.

Usage:
  python import_questions.py [--file out/kimya_dim_yoxlama.xlsx] [--out out/import_kimya.sql] [--check]
"""
import argparse
import collections
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from build_xlsx import COLUMNS

HERE = Path(__file__).parent
SUBJECT_SLUG = "kimya"
LOADABLE_TYPES = ("qapali", "aciq_eded")
SECTION_ORDER = ["Ümumi kimya", "Qeyri-üzvi kimya", "Üzvi kimya"]
NUMERIC = re.compile(r"-?\d+(?:[.,]\d+)?")


class ImportErrors(Exception):
    pass


def lit(value):
    """SQL string literal (standard_conforming_strings is on, so only quotes need escaping)."""
    return "NULL::text" if value is None or value == "" else "'" + str(value).replace("'", "''") + "'"


def read_rows(path):
    ws = load_workbook(path)["suallar"]
    header = [c.value for c in ws[1]]
    if header != COLUMNS:
        raise ImportErrors([f"columns differ from the expected template: {header}"])
    return [
        (i, dict(zip(header, ("" if v is None else str(v).strip() for v in row))))
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
    ]


def select_rows(rows):
    """Returns (rows to load, Counter of skip reasons)."""
    skipped, chosen = collections.Counter(), []
    for i, r in rows:
        if r["aktiv"] != "bəli":
            skipped["aktiv deyil (müəllim təsdiqləməyib)"] += 1
        elif r["tip"] not in LOADABLE_TYPES:
            skipped[f"tip {r['tip']!r} ilk versiyaya daxil deyil"] += 1
        elif r["çətinlik"] == "":
            skipped["çətinlik doldurulmayıb"] += 1
        else:
            chosen.append((i, r))
    return chosen, skipped


def validate(chosen, images_dir):
    errors, seen, section_of = [], {}, {}
    for i, r in chosen:
        def err(msg):
            errors.append(f"sətir {i} ({r['kod']}): {msg}")

        if not re.fullmatch(r"KIM-\d{4}", r["kod"]):
            err("kod KIM-0000 formatında deyil")
        if r["kod"] in seen:
            err(f"kod təkrarlanır (sətir {seen[r['kod']]})")
        seen[r["kod"]] = i
        if not re.fullmatch(r"[1-5]", r["çətinlik"]):
            err(f"çətinlik 1-5 olmalıdır, yazılıb: {r['çətinlik']!r}")
        if not r["bölmə"] or not r["mövzu"]:
            err("bölmə və ya mövzu boşdur")
        elif section_of.setdefault(r["mövzu"], r["bölmə"]) != r["bölmə"]:
            err(f"mövzu {r['mövzu']!r} iki fərqli bölmədə")
        if not r["sual mətni"]:
            err("sual mətni boşdur")
        if r["səviyyə"] != "qebul":
            err(f"səviyyə qebul olmalıdır: {r['səviyyə']!r}")
        if not re.fullmatch(r"[1-5](,[1-5])*", r["qruplar"]):
            err(f"qruplar düzgün deyil: {r['qruplar']!r}")
        if not re.fullmatch(r"\d{4}", r["il"]):
            err(f"il tək 4 rəqəmli olmalıdır: {r['il']!r}")
        if not r["mənbə"] or not r["mənbədə yeri"]:
            err("mənbə və ya mənbədə yeri boşdur")
        if r["tip"] == "qapali":
            missing = [L for L in "ABCDE" if not r[L]]
            if missing:
                err(f"variantlar boşdur: {missing}")
            if r["düzgün cavab"] not in list("ABCDE"):
                err(f"düzgün cavab A-E olmalıdır: {r['düzgün cavab']!r}")
        elif not NUMERIC.fullmatch(r["düzgün cavab"]):
            err(f"aciq_eded cavabı ədəd olmalıdır (baza onu ədədə çevirir): {r['düzgün cavab']!r}")
        for col in ("sual mətni", "A", "B", "C", "D", "E", "düzgün cavab", "həll"):
            v = r[col]
            if r"\ce" in v:
                err(f"{col}: \\ce istifadə olunub")
            if v.count("$") % 2:
                err(f"{col}: $ cüt deyil")
        if r["şəkil"] and not (images_dir / r["şəkil"]).exists():
            err(f"şəkil faylı yoxdur: {r['şəkil']}")
    if errors:
        raise ImportErrors(errors)


def build_sql(chosen):
    topics = {}
    for _i, r in chosen:
        topics[r["mövzu"]] = r["bölmə"]
    sections = [s for s in SECTION_ORDER if s in set(topics.values())]
    sections += sorted(set(topics.values()) - set(sections))
    topic_rows = sorted(topics.items(), key=lambda kv: (sections.index(kv[1]), kv[0]))

    q_rows = []
    for _i, r in chosen:
        options = json.dumps({L: r[L] for L in "ABCDE"}, ensure_ascii=False) if r["tip"] == "qapali" else None
        groups = "{" + r["qruplar"] + "}"
        q_rows.append(
            "  (" + ", ".join([
                lit(r["kod"]), lit(r["bölmə"]), lit(r["mövzu"]), lit(r["tip"]), lit(r["sual mətni"]),
                lit(r["şəkil"]), lit(options), lit(r["düzgün cavab"]), lit(r["çətinlik"]), lit(r["həll"]),
                lit(groups), lit(r["il"]), lit(r["mənbədə yeri"]),
            ]) + ")"
        )
    codes = ", ".join(lit(r["kod"]) for _i, r in chosen)
    section_values = ",\n".join(f"  ({lit(s)}, {n})" for n, s in enumerate(sections, 1))
    topic_values = ",\n".join(f"  ({lit(sec)}, {lit(name)}, {n})" for n, (name, sec) in enumerate(topic_rows, 1))

    return f"""-- Generated by tools/dim/import_questions.py. {len(chosen)} approved chemistry questions.
-- Run once in Supabase SQL Editor. Safe to re-run: topics only if missing, questions upserted by code.
set standard_conforming_strings = on;
begin;

-- 0) pre-flight: stop with a clear message instead of failing half way
do $$
begin
  if not exists (select 1 from subjects where slug = '{SUBJECT_SLUG}') then
    raise exception 'subjects cədvəlində slug={SUBJECT_SLUG} yoxdur';
  end if;
  if not exists (select 1 from pg_enum e join pg_type t on t.oid = e.enumtypid
                  where t.typname = 'exam_level' and e.enumlabel = 'qebul') then
    raise exception 'exam_level enum-unda qebul etiketi yoxdur: import dayandırıldı';
  end if;
  if (select count(*) from pg_enum e join pg_type t on t.oid = e.enumtypid
       where t.typname = 'question_type' and e.enumlabel in ('qapali', 'aciq_eded')) < 2 then
    raise exception 'question_type enum-unda qapali/aciq_eded etiketləri yoxdur: import dayandırıldı';
  end if;
end $$;

-- 1) sections (top-level topics)
insert into topics (subject_id, parent_id, name, sort_order)
select s.id, null, v.name, v.ord
  from subjects s,
       (values
{section_values}
       ) as v(name, ord)
 where s.slug = '{SUBJECT_SLUG}'
on conflict (subject_id, parent_id, name) do nothing;

-- 2) topics under their section
insert into topics (subject_id, parent_id, name, sort_order)
select s.id, p.id, v.name, v.ord
  from subjects s,
       (values
{topic_values}
       ) as v(section, name, ord)
  join topics p on p.parent_id is null and p.name = v.section
 where s.slug = '{SUBJECT_SLUG}' and p.subject_id = s.id
on conflict (subject_id, parent_id, name) do nothing;

-- 3) questions, upserted by code
insert into questions (code, topic_id, type, body, image_path, options, correct_answer, difficulty, solution,
                       exam_levels, groups, year, source, source_ref, is_active)
select v.code, t.id, v.type::question_type, v.body, v.image_path, v.options::jsonb, v.correct,
       v.difficulty::smallint, v.solution, array['qebul']::exam_level[], v.groups::smallint[],
       v.year::smallint, 'DİM', v.source_ref, true
  from (values
{",\n".join(q_rows)}
       ) as v(code, section, topic, type, body, image_path, options, correct, difficulty, solution,
              groups, year, source_ref)
  join subjects s on s.slug = '{SUBJECT_SLUG}'
  join topics sec on sec.subject_id = s.id and sec.parent_id is null and sec.name = v.section
  join topics t on t.subject_id = s.id and t.parent_id = sec.id and t.name = v.topic
on conflict (code) do update set
  topic_id = excluded.topic_id, type = excluded.type, body = excluded.body, image_path = excluded.image_path,
  options = excluded.options, correct_answer = excluded.correct_answer, difficulty = excluded.difficulty,
  solution = excluded.solution, exam_levels = excluded.exam_levels, groups = excluded.groups,
  year = excluded.year, source = excluded.source, source_ref = excluded.source_ref,
  is_active = excluded.is_active;

-- 4) every approved question must now exist (a missing topic would silently drop a row from the join)
do $$
declare n integer;
begin
  select count(*) into n from questions where code in ({codes});
  if n <> {len(chosen)} then
    raise exception 'Gözlənilən {len(chosen)} sual, bazada % tapıldı: geri alınır', n;
  end if;
end $$;

commit;

-- 5) result
select s.name as fenn, count(*) as sual_sayi
  from questions q
  join topics t on t.id = q.topic_id
  join subjects s on s.id = t.subject_id
 where s.slug = '{SUBJECT_SLUG}' and q.source = 'DİM'
 group by s.name;
"""


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", type=Path, default=HERE / "out" / "kimya_dim_yoxlama.xlsx")
    ap.add_argument("--out", type=Path, default=HERE / "out" / "import_kimya.sql")
    ap.add_argument("--check", action="store_true", help="validate only, write nothing")
    args = ap.parse_args(argv)

    rows = read_rows(args.file)
    chosen, skipped = select_rows(rows)
    print(f"Excel-də {len(rows)} sual: yüklənəcək {len(chosen)}, buraxılır {sum(skipped.values())}")
    for reason, n in skipped.most_common():
        print(f"  buraxıldı: {n} × {reason}")
    try:
        validate(chosen, HERE / "images")
    except ImportErrors as e:
        print(f"\nXƏTA: {len(e.args[0])} problem tapıldı, heç nə yazılmadı:")
        for line in e.args[0]:
            print("  -", line)
        return 1
    if not chosen:
        print("Yüklənəcək sual yoxdur (müəllim hələ təsdiqləməyib).")
        return 0
    if args.check:
        print("Yoxlama uğurludur (--check: fayl yazılmadı).")
        return 0
    args.out.write_text(build_sql(chosen), encoding="utf-8")
    print(f"SQL yazıldı: {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
