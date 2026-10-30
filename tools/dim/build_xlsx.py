"""Builds out/kimya_dim_yoxlama.xlsx from out/parsed.json (run extract_kimya.py first).

Usage: python build_xlsx.py
"""
import collections
import json
import re
import sys
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from chem_text import MARKS, to_tex
from topic_map import map_topic
from extract_kimya import FILES, IMAGES, OUT, RAW, collapse, crop_figure, parse_options, plain_key

COLUMNS = [
    "kod", "fənn", "bölmə", "mövzu", "DİM mövzusu", "tip", "sual mətni", "A", "B", "C", "D", "E", "düzgün cavab", "çətinlik", "həll",
    "şəkil", "səviyyə", "qruplar", "il", "mənbə", "mənbədə yeri", "aktiv", "yoxlanmalı", "qeyd", "sinif",
    "izah şəkli",
]
ROMAN_CLASS = {"VII": "7", "VIII": "8", "IX": "9", "X": "10", "XI": "11"}
NUMERIC = re.compile(r"-?\d+(?:[.,]\d+)?")
YELLOW = PatternFill("solid", fgColor="FFF2CC")


def loose_key(s):
    """Answer key without 'X-', 'a =' style labels, so 'X-HCOOH Y-CH3OH' equals the option 'HCOOH CH3OH'."""
    s = re.sub(f"[{MARKS}]", "", s)
    s = re.sub(r"(?<![A-Za-z0-9])[A-Za-z]\s*[:=–-]+\s*", "", s)
    return re.sub(r"[\s.;,:=]", "", s).replace("–", "-").replace("−", "-").lower()


def convert(raw, flags):
    return to_tex(collapse(raw), flags)


def build_row(r, conv=convert):
    """[conv] turns PDF text into the stored text; chemistry uses LaTeX, other subjects plain text."""
    flags, reasons = set(), []
    stem_raw = collapse(f"{r['ctx_text']} {r['q_text']}")
    stem_raw = re.sub(r"^\d+\.\s*", "", stem_raw)
    stem_raw, opts_raw = parse_options(stem_raw)

    if r["unknown"]:
        reasons.append("naməlum simvol " + ", ".join(r["unknown"]))
    reasons += r.get("tags", [])
    reasons += r["visual"]
    reasons += [f"izahda: {v}" for v in r.get("expl_visual", [])]
    if r["short_lines"] >= 5 and "cədvəl" not in r["visual"]:
        reasons.append("cədvəl/parçalanmış mətn")
    if r["ctx_text"]:
        reasons.append("situasiya tapşırığı (ortaq şərt + qrafik)")

    answer_raw = r["answer"]
    if opts_raw:
        qtype = "qapali"
        letters = [L for L, t in opts_raw.items() if answer_raw and plain_key(t) == plain_key(answer_raw)]
        if answer_raw and answer_raw.strip() in opts_raw:
            letters = [answer_raw.strip()]
        if not letters and answer_raw:
            letters = [L for L, t in opts_raw.items() if loose_key(t) == loose_key(answer_raw)]
        if len(letters) == 1:
            correct = letters[0]
        else:
            correct = ""
            reasons.append("düzgün cavab variantlarla uyğunlaşmadı" if answer_raw else "düzgün cavab yoxdur")
    else:
        correct = ""
        if answer_raw is None:
            qtype = "yazili"
            reasons.append("düzgün cavab yoxdur")
        elif NUMERIC.fullmatch(re.sub(f"[{MARKS}]", "", answer_raw).strip()):
            qtype, correct = "aciq_eded", answer_raw.strip()
        else:
            qtype, correct = "yazili", conv(answer_raw, flags)
            reasons.append("yazılı/uyğunluq cavabı: əl ilə yoxla")
    if not r["topic"]:
        reasons.append("mövzu yoxdur")

    stem = conv(stem_raw, flags)
    options = {L: conv(t, flags) for L, t in opts_raw.items()}
    solution = conv(r["expl"], flags)
    for f in sorted(flags):
        reasons.append({"formul": "formul tam çevrilmədi", "yük": "ion yükü", "rabitə": "rabitə (=, ≡)"}[f])

    cls = r["class"]
    return {
        "dedup": plain_key(stem_raw + "".join(opts_raw.values())),
        "topic": re.sub(f"[{MARKS}]", "", r["topic"]).strip(),
        "type": qtype, "stem": stem, "options": options, "correct": correct, "solution": solution,
        "reasons": list(dict.fromkeys(reasons)),
        "sinif": ", ".join(ROMAN_CLASS.get(c.strip(), c.strip()) for c in cls.split(",")) if cls else "",
        "groups": [r["group"]], "years": [r["year"]],
        "places": [f"{r['exam']}, A variantı №{r['a_no']}"],
        "crop": (r["file"], r["q_lines"]) if (r["visual"] or r["ctx_text"] or r["short_lines"] >= 5) else None,
        "file": r["file"],
        "expl_crop": r.get("expl_q_lines") or None,
    }


def excel_row(kod, subject_name, b, section, topic, image, expl_image, reasons):
    """One Excel row (in COLUMNS order) from a [build_row] result."""
    opts = b["options"]
    return [
        kod, subject_name, section, topic, b["topic"], b["type"], b["stem"],
        opts.get("A", ""), opts.get("B", ""), opts.get("C", ""), opts.get("D", ""), opts.get("E", ""),
        b["correct"], "", b["solution"], image, "qebul",
        ",".join(dict.fromkeys(b["groups"])), ",".join(str(y) for y in dict.fromkeys(b["years"])),
        "DİM", "; ".join(b["places"]), "xeyr", "bəli" if reasons else "", "; ".join(reasons), b["sinif"],
        expl_image,
    ]


def main():
    target = OUT / "kimya_dim_yoxlama.xlsx"
    if "--out" in sys.argv:
        target = Path(sys.argv[sys.argv.index("--out") + 1])
    if target.exists() and "--force" not in sys.argv:
        sys.exit(
            f"{target.name} already exists. Codes (KIM-xxxx) are keys in the database and the teacher edits this file: "
            "rebuilding would renumber them. Pass --force only before any review or upload."
        )
    rows = json.loads((OUT / "parsed.json").read_text(encoding="utf-8"))
    merged, seen, per_file = [], {}, collections.Counter()
    for r in rows:
        per_file[r["file"]] += 1
        b = build_row(r)
        if b["dedup"] in seen:
            first = seen[b["dedup"]]
            first["groups"] += b["groups"]
            first["years"] += b["years"]
            first["places"] += b["places"]
            continue
        seen[b["dedup"]] = b
        merged.append(b)

    for old in IMAGES.glob("KIM-*.png"):
        old.unlink()
    docs = {stem: pymupdf.open(RAW / f"{stem}.pdf") for stem, *_ in FILES}
    wb = Workbook()
    ws = wb.active
    ws.title = "suallar"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    flagged_by_file = collections.Counter()
    topic_counts = collections.Counter()
    for n, b in enumerate(merged, 1):
        kod = f"KIM-{n:04d}"
        image = ""
        if b["crop"]:
            crop_figure(docs[b["crop"][0]], b["crop"][1], kod)
            image = f"{kod}.png"
        expl_image = ""
        if b["expl_crop"]:
            expl_image = f"{kod}_izah.png"
            crop_figure(docs[b["file"]], b["expl_crop"], f"{kod}_izah", pad=0)
        flagged = bool(b["reasons"])
        if flagged:
            flagged_by_file[b["file"]] += 1
        topic_counts[b["topic"] or "(mövzu yoxdur)"] += 1
        opts = b["options"]
        section, topic, _note = map_topic(b["topic"])
        ws.append([
            kod, "Kimya", section, topic, b["topic"], b["type"], b["stem"],
            opts.get("A", ""), opts.get("B", ""), opts.get("C", ""), opts.get("D", ""), opts.get("E", ""),
            b["correct"], "", b["solution"], image, "qebul",
            ",".join(dict.fromkeys(b["groups"])), ",".join(str(y) for y in dict.fromkeys(b["years"])),
            "DİM", "; ".join(b["places"]), "xeyr", "bəli" if flagged else "", "; ".join(b["reasons"]), b["sinif"],
            expl_image,
        ])
        if flagged:
            for c in ws[ws.max_row]:
                c.fill = YELLOW
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVWXYZ", (10, 8, 16, 34, 34, 10, 60, 24, 24, 24, 24, 24, 12, 9, 50, 14, 9, 8, 8, 6, 34, 7, 11, 40, 7, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    s = wb.create_sheet("xülasə")
    s.append(["Fayl", "İmtahan", "Çıxarılan sual", "Bu faylda işarələnmiş (dublikatsız)"])
    for stem, exam, _g, _y in FILES:
        s.append([stem, exam, per_file[stem], flagged_by_file[stem]])
    s.append([])
    s.append(["Cəmi çıxarılan", "", sum(per_file.values()), ""])
    s.append(["Dublikatsız sual", "", len(merged), sum(flagged_by_file.values())])
    s.append([])
    s.append(["Mövzu", "", "Sual sayı"])
    for t, c in sorted(topic_counts.items(), key=lambda kv: (-kv[1], kv[0])):
        s.append([t, "", c])
    s.column_dimensions["A"].width = 60
    s.column_dimensions["B"].width = 30
    s.column_dimensions["C"].width = 16
    s.column_dimensions["D"].width = 34
    wb.save(target)
    print(f"questions {len(merged)} (from {sum(per_file.values())}), flagged {sum(flagged_by_file.values())}")
    print("reasons:", collections.Counter(r.split(" ")[0] + " " + (r.split(" ")[1] if " " in r else "") for b in merged for r in b["reasons"]).most_common(12))


if __name__ == "__main__":
    main()
