"""Appends questions that are not yet in the subject's Excel (new PDFs added to extract_*.py FILES).

Existing rows, codes, teacher decisions and images stay untouched. Only:
  - brand-new unique questions are added with the next free code (flagged for review, aktiv = xeyr);
  - a question already in the Excel that also appears in a new exam gets that exam added to
    "qruplar", "il" and "mənbədə yeri" (so it is never stored twice).
Questions that are only similar to another one get an "oxşar sual" note.
A backup is written first (out/<fənn>_dim_yoxlama.backup.xlsx, overwritten each time).

Usage: python append_new.py --subject kimya|biologiya [--dry-run]
"""
import argparse
import difflib
import json
import shutil
import sys

import pymupdf
from openpyxl import load_workbook

from build_xlsx import COLUMNS, YELLOW, build_row, convert, excel_row
from build_biologiya import plain
from extract_biologiya import BIO_FILES
from extract_kimya import FILES, IMAGES, RAW, crop_figure
from subjects import SUBJECTS
from topic_map import map_topic as map_kimya
from topic_map_biologiya import map_topic as map_biologiya

SIMILAR = 0.9
SETUP = {
    "kimya": {"files": FILES, "parsed": "parsed.json", "conv": convert, "map": map_kimya},
    "biologiya": {"files": BIO_FILES, "parsed": "parsed_biologiya.json", "conv": plain, "map": map_biologiya},
}


def merge_all(rows, conv):
    """Same exact-duplicate merge as the build scripts, in the same order, so old codes keep their position."""
    merged, seen = [], {}
    for r in rows:
        b = build_row(r, conv)
        first = seen.get(b["dedup"])
        if first is None:
            seen[b["dedup"]] = b
            merged.append(b)
        else:
            for key in ("groups", "years", "places"):
                first[key] += b[key]
    return merged


def similar_to(index, merged):
    a = merged[index]["dedup"]
    out = []
    for j, other in enumerate(merged):
        if j == index:
            continue
        sm = difflib.SequenceMatcher(None, a, other["dedup"])
        if sm.real_quick_ratio() >= SIMILAR and sm.quick_ratio() >= SIMILAR and sm.ratio() >= SIMILAR:
            out.append(j)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--subject", choices=SUBJECTS, required=True)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    cfg, setup = SUBJECTS[args.subject], SETUP[args.subject]
    path = cfg["excel"]

    rows = json.loads((path.parent / setup["parsed"]).read_text(encoding="utf-8"))
    merged = merge_all(rows, setup["conv"])
    codes = [f"{cfg['prefix']}-{n:04d}" for n in range(1, len(merged) + 1)]

    wb = load_workbook(path)
    ws = wb["suallar"]
    header = [c.value for c in ws[1]]
    if header != COLUMNS:
        sys.exit(f"columns differ from the template: {header}")
    col = {h: i + 1 for i, h in enumerate(header)}
    existing = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    n_old = len(existing)

    # the first n_old merged rows must be exactly the Excel rows; otherwise the old order changed and codes would shift
    for kod, b in zip(codes[:n_old], merged[:n_old]):
        if kod not in existing or ws.cell(existing[kod], col["sual mətni"]).value != b["stem"]:
            sys.exit(f"{kod}: Excel text differs from the parsed PDF order; refusing to append (codes would shift)")

    updated = 0
    for kod, b in zip(codes[:n_old], merged[:n_old]):
        r = existing[kod]
        new = {
            "qruplar": ",".join(dict.fromkeys(b["groups"])),
            "il": ",".join(str(y) for y in dict.fromkeys(b["years"])),
            "mənbədə yeri": "; ".join(b["places"]),
        }
        if any(str(ws.cell(r, col[k]).value or "") != v for k, v in new.items()):
            updated += 1
            if not args.dry_run:
                for k, v in new.items():
                    ws.cell(r, col[k]).value = v

    added = merged[n_old:]
    print(f"{args.subject}: Excel has {n_old}, parsed unique {len(merged)} -> {len(added)} new, {updated} existing get extra exams")
    if args.dry_run or not added and not updated:
        return
    shutil.copy(path, path.with_suffix(".backup.xlsx"))

    docs = {stem: pymupdf.open(RAW / f"{stem}.pdf") for stem, *_ in setup["files"]}
    flagged = 0
    for i, b in enumerate(added, start=n_old):
        kod = codes[i]
        image = expl_image = ""
        if b["crop"]:
            crop_figure(docs[b["crop"][0]], b["crop"][1], kod)
            image = f"{kod}.png"
        if b["expl_crop"]:
            expl_image = f"{kod}_izah.png"
            crop_figure(docs[b["file"]], b["expl_crop"], f"{kod}_izah", pad=0)
        section, topic, note = setup["map"](b["topic"])
        reasons = list(b["reasons"])
        if note:
            reasons.append(note)
        near = similar_to(i, merged)
        if near:
            reasons.append("oxşar sual: " + ", ".join(codes[j] for j in near))
        ws.append(excel_row(kod, cfg["name"], b, section, topic, image, expl_image, reasons))
        if reasons:
            flagged += 1
            for c in ws[ws.max_row]:
                c.fill = YELLOW
    wb.save(path)
    print(f"added {len(added)} rows ({flagged} flagged for review); backup: {path.with_suffix('.backup.xlsx').name}")


if __name__ == "__main__":
    main()
