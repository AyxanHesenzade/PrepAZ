"""Validates out/kimya_dim_yoxlama.xlsx. Exit code 1 on errors.

The original tools/import_questions.py (--check) was lost when the old prepaz/ folder was deleted,
so this checker enforces the rules from TASK_01 itself.

Usage: python check.py [path.xlsx]
"""
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from build_xlsx import COLUMNS

DEFAULT = Path(__file__).parent / "out" / "kimya_dim_yoxlama.xlsx"
TYPES = {"qapali", "aciq_eded", "yazili"}
TEXT_COLS = ["sual mətni", "A", "B", "C", "D", "E", "düzgün cavab", "həll"]


def check(path):
    ws = load_workbook(path)["suallar"]
    header = [c.value for c in ws[1]]
    errors, warnings = [], []
    if header != COLUMNS:
        errors.append(f"columns differ: {header}")
        return errors, warnings
    seen, empty_difficulty = set(), 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(COLUMNS, ("" if v is None else str(v) for v in row)))
        k = r["kod"]
        flagged = r["yoxlanmalı"] == "bəli"

        def err(msg):
            errors.append(f"{k}: {msg}")

        if not re.fullmatch(r"KIM-\d{4}", k) or k in seen:
            err("bad or duplicate code")
        seen.add(k)
        if r["fənn"] != "Kimya":
            err("fənn must be Kimya")
        if r["tip"] not in TYPES:
            err(f"bad type {r['tip']!r}")
        if r["mövzu"] and not r["bölmə"]:
            err("mövzu without bölmə")
        if not r["sual mətni"]:
            err("empty question text")
        if not r["mənbə"] or not r["mənbədə yeri"]:
            err("missing source")
        if r["səviyyə"] != "qebul":
            err("səviyyə must be qebul")
        if r["aktiv"] not in ("xeyr", "bəli"):
            err("aktiv must be bəli or xeyr")
        if r["aktiv"] == "bəli":
            if r["çətinlik"] not in ("1", "2", "3", "4", "5"):
                err("approved question needs çətinlik 1-5")
            if r["tip"] not in ("qapali", "aciq_eded"):
                err("only qapali/aciq_eded can be approved in this version")
        if not re.fullmatch(r"[14](,[14])*", r["qruplar"]) or not re.fullmatch(r"\d{4}(,\d{4})*", r["il"]):
            err("bad qruplar/il")
        if r["çətinlik"] == "":
            empty_difficulty += 1
        for col in TEXT_COLS:
            v = r[col]
            if r"\ce" in v:
                err(f"\\ce used in {col}")
            if v.count("$") % 2:
                err(f"unbalanced $ in {col}")
            if v.count("{") != v.count("}"):
                err(f"unbalanced braces in {col}")
            if any(not m.startswith("\\mathrm{") for m in re.findall(r"\$([^$]*)\$", v)):
                err(f"math not in $\\mathrm{{...}}$ form in {col}")
        if r["tip"] == "qapali":
            if any(not r[L] for L in "ABCDE"):
                err("closed question needs A-E")
            if r["düzgün cavab"] not in list("ABCDE") and not flagged:
                err("no answer letter and not flagged")
        elif r["tip"] == "aciq_eded" and not re.fullmatch(r"-?\d+([.,]\d+)?", r["düzgün cavab"]):
            err("aciq_eded answer is not a number")
        if not r["düzgün cavab"] and not flagged:
            err("empty answer and not flagged")
        if r["şəkil"] and not (Path(path).parent.parent / "images" / r["şəkil"]).exists():
            err("image file missing")
    if empty_difficulty:
        warnings.append(f"{empty_difficulty} rows with empty çətinlik (expected: teacher fills it)")
    return errors, warnings


if __name__ == "__main__":
    errs, warns = check(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT)
    for w in warns:
        print("WARN", w)
    for e in errs:
        print("ERROR", e)
    print(f"{len(errs)} errors, {len(warns)} warnings")
    sys.exit(1 if errs else 0)
